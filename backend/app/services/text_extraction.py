"""
Text Extraction Service
Extracts text from PDF and DOCX files
"""

import pdfplumber
import PyPDF2
import docx
from typing import Dict, Optional
import os


def extract_text_from_pdf(file_path: str) -> Dict[str, any]:
    """
    Extract text from PDF file using pdfplumber with PyPDF2 fallback
    
    Args:
        file_path: Path to PDF file
        
    Returns:
        Dictionary with extracted text and metadata
    """
    result = {
        "text": "",
        "pages": 0,
        "method": "pdfplumber",
        "success": False,
        "error": None
    }
    
    try:
        # Try pdfplumber first (better layout preservation)
        with pdfplumber.open(file_path) as pdf:
            result["pages"] = len(pdf.pages)
            text_parts = []
            
            for page in pdf.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
                    
                # Extract tables if present
                tables = page.extract_tables()
                for table in tables:
                    for row in table:
                        if row:
                            text_parts.append(' | '.join(str(cell) for cell in row if cell))
            
            result["text"] = '\n'.join(text_parts)
            result["success"] = True
            
    except Exception as e:
        # Fallback to PyPDF2
        try:
            result["method"] = "PyPDF2"
            with open(file_path, 'rb') as file:
                pdf_reader = PyPDF2.PdfReader(file)
                result["pages"] = len(pdf_reader.pages)
                text_parts = []
                
                for page in pdf_reader.pages:
                    text = page.extract_text()
                    if text:
                        text_parts.append(text)
                
                result["text"] = '\n'.join(text_parts)
                result["success"] = True
                
        except Exception as fallback_error:
            result["error"] = f"pdfplumber: {str(e)}, PyPDF2: {str(fallback_error)}"
            result["success"] = False
    
    return result


def extract_text_from_docx(file_path: str) -> Dict[str, any]:
    """
    Extract text from DOCX file
    
    Args:
        file_path: Path to DOCX file
        
    Returns:
        Dictionary with extracted text and metadata
    """
    result = {
        "text": "",
        "paragraphs": 0,
        "success": False,
        "error": None
    }
    
    try:
        doc = docx.Document(file_path)
        text_parts = []
        
        # Extract paragraphs
        for para in doc.paragraphs:
            if para.text.strip():
                text_parts.append(para.text)
        
        # Extract tables
        for table in doc.tables:
            for row in table.rows:
                row_text = ' | '.join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    text_parts.append(row_text)
        
        result["text"] = '\n'.join(text_parts)
        result["paragraphs"] = len(doc.paragraphs)
        result["success"] = True
        
    except Exception as e:
        result["error"] = str(e)
        result["success"] = False
    
    return result


def extract_text_from_file(file_path: str) -> Dict[str, any]:
    """
    Extract text from file (auto-detect PDF or DOCX)
    
    Args:
        file_path: Path to file
        
    Returns:
        Dictionary with extracted text and metadata
    """
    if not os.path.exists(file_path):
        return {
            "text": "",
            "success": False,
            "error": "File not found"
        }
    
    file_ext = file_path.rsplit('.', 1)[-1].lower()
    
    if file_ext == 'pdf':
        return extract_text_from_pdf(file_path)
    elif file_ext in ['docx', 'doc']:
        return extract_text_from_docx(file_path)
    else:
        return {
            "text": "",
            "success": False,
            "error": f"Unsupported file type: {file_ext}"
        }


def detect_layout(text: str) -> Dict[str, any]:
    """
    Analyze document layout and structure
    
    Args:
        text: Extracted text
        
    Returns:
        Dictionary with layout information
    """
    lines = text.split('\n')
    
    return {
        "total_lines": len(lines),
        "non_empty_lines": len([l for l in lines if l.strip()]),
        "avg_line_length": sum(len(l) for l in lines) / len(lines) if lines else 0,
        "has_tables": '|' in text,  # Simple table detection
        "estimated_sections": len([l for l in lines if l.isupper() and len(l.split()) <= 5])
    }
